"""Generate a formatted Word (.docx) report from the analysis results."""

import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# ── Colour scheme ──────────────────────────────────────────────────────────────
_GREEN = RGBColor(0x16, 0xa3, 0x4a)
_AMBER = RGBColor(0xd9, 0x77, 0x06)
_RED = RGBColor(0xdc, 0x26, 0x26)
_DARK_BLUE_HEX = "1F3864"
_LIGHT_BLUE_HEX = "DBEAFE"
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _compliance_color(pct: int) -> RGBColor:
    if pct >= 80:
        return _GREEN
    if pct >= 60:
        return _AMBER
    return _RED


def _compliance_label(pct: int) -> str:
    if pct >= 80:
        return "ALTO CUMPLIMIENTO"
    if pct >= 60:
        return "CUMPLIMIENTO MODERADO"
    if pct >= 40:
        return "CUMPLIMIENTO BAJO"
    return "CUMPLIMIENTO INSUFICIENTE"


def _bar(pct: int, width: int = 25) -> str:
    filled = round(pct / 100 * width)
    return "█" * filled + "░" * (width - filled)


# ── XML helpers ────────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C0C0C0")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Paragraph helpers ──────────────────────────────────────────────────────────

def _colored_run(para, text: str, color: RGBColor, bold: bool = False, size_pt: int = 11) -> None:
    run = para.add_run(text)
    run.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size_pt)


def _bullet_runs(doc: Document, items: list[str], color: RGBColor) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        _colored_run(p, item, color)


# ── Main generator ─────────────────────────────────────────────────────────────

def generate_word_report(
    analysis: dict[str, Any],
    eval_filename: str,
    output_path: Path,
) -> None:
    doc = Document()

    # Global font defaults
    for style_name in ("Normal", "List Bullet", "List Number"):
        try:
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style.font.size = Pt(11)
        except KeyError:
            pass

    overall = int(analysis.get("overall_compliance", 0))
    color = _compliance_color(overall)
    label = _compliance_label(overall)
    sections: list[dict] = analysis.get("sections", [])
    now = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")

    # ── Cover / header ──────────────────────────────────────────────────────
    title = doc.add_heading("REPORTE DE VALIDACIÓN DE DOCUMENTO ACADÉMICO", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Generado con Inteligencia Artificial · Claude Opus 4").italic = True

    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.add_run("Documento evaluado: ").bold = True
    meta.add_run(eval_filename)
    meta2 = doc.add_paragraph()
    meta2.add_run("Fecha de evaluación: ").bold = True
    meta2.add_run(now)

    _add_horizontal_rule(doc)

    # ── Overall compliance ──────────────────────────────────────────────────
    doc.add_heading("RESUMEN EJECUTIVO", 1)

    score_para = doc.add_paragraph()
    score_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _colored_run(score_para, f"{overall}%", color, bold=True, size_pt=48)

    label_para = doc.add_paragraph()
    label_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _colored_run(label_para, label, color, bold=True, size_pt=14)

    bar_para = doc.add_paragraph()
    bar_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _colored_run(bar_para, _bar(overall), color, size_pt=14)

    doc.add_paragraph()

    exec_summary = analysis.get("executive_summary", "")
    if exec_summary:
        sum_para = doc.add_paragraph(exec_summary)
        sum_para.paragraph_format.space_after = Pt(12)
        sum_para.paragraph_format.left_indent = Inches(0.3)

    doc.add_page_break()

    # ── Section-by-section analysis ─────────────────────────────────────────
    doc.add_heading("ANÁLISIS POR SECCIONES", 1)

    for idx, section in enumerate(sections, 1):
        sec_pct = int(section.get("compliance_percentage", 0))
        sec_color = _compliance_color(sec_pct)
        sec_label = _compliance_label(sec_pct)

        heading = doc.add_heading(
            f"{idx}. {section.get('section_name', f'Sección {idx}')}", 2
        )

        # Compliance bar for this section
        bar_p = doc.add_paragraph()
        _colored_run(bar_p, f"{_bar(sec_pct, 20)}  {sec_pct}% — {sec_label}", sec_color, bold=True)

        # Observations
        doc.add_heading("Observaciones", 3)
        obs = doc.add_paragraph(section.get("observations", "—"))
        obs.paragraph_format.left_indent = Inches(0.2)

        # Strengths
        strengths = section.get("strengths", [])
        if strengths:
            doc.add_heading("Fortalezas", 3)
            _bullet_runs(doc, strengths, _GREEN)

        # Corrections
        corrections = section.get("corrections", [])
        if corrections:
            doc.add_heading("Correcciones requeridas", 3)
            _bullet_runs(doc, corrections, _RED)

        _add_horizontal_rule(doc)
        doc.add_paragraph()

    doc.add_page_break()

    # ── Recommendations ─────────────────────────────────────────────────────
    doc.add_heading("RECOMENDACIONES GENERALES", 1)
    recs = analysis.get("general_recommendations", [])
    if recs:
        for rec in recs:
            p = doc.add_paragraph(style="List Number")
            p.add_run(rec)
    else:
        doc.add_paragraph("Sin recomendaciones adicionales.")

    doc.add_paragraph()

    # ── Missing elements ────────────────────────────────────────────────────
    missing = analysis.get("missing_elements", [])
    if missing:
        doc.add_heading("ELEMENTOS FALTANTES O INCOMPLETOS", 1)
        _bullet_runs(doc, missing, _RED)
        doc.add_paragraph()

    doc.add_page_break()

    # ── Integrated comments table ───────────────────────────────────────────
    doc.add_heading("COMENTARIOS INTEGRADOS POR SECCIÓN", 1)
    intro = doc.add_paragraph(
        "La siguiente tabla resume las correcciones específicas identificadas "
        "por sección, al estilo de comentarios de revisión académica."
    )
    intro.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    # Header row
    hdr = table.rows[0].cells
    headers = ["SECCIÓN", "CUMPLIMIENTO", "COMENTARIOS DEL REVISOR"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = _WHITE
        run.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_bg(cell, _DARK_BLUE_HEX)

    # Data rows
    for section in sections:
        sec_pct = int(section.get("compliance_percentage", 0))
        row = table.add_row().cells

        # Col 0 – section name
        row[0].text = section.get("section_name", "—")
        row[0].paragraphs[0].runs[0].font.size = Pt(10)

        # Col 1 – percentage with colour
        pct_para = row[1].paragraphs[0]
        pct_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _colored_run(pct_para, f"{sec_pct}%", _compliance_color(sec_pct), bold=True, size_pt=12)
        _set_cell_bg(row[1], _LIGHT_BLUE_HEX)

        # Col 2 – corrections or OK note
        corrections = section.get("corrections", [])
        if corrections:
            comment_text = "\n• ".join([""] + corrections).lstrip("\n")
        else:
            comment_text = "✓ Sin correcciones requeridas."
        row[2].text = comment_text
        row[2].paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # ── Footer note ─────────────────────────────────────────────────────────
    _add_horizontal_rule(doc)
    footer = doc.add_paragraph(
        "Este reporte fue generado automáticamente mediante análisis de inteligencia artificial. "
        "Se recomienda complementar con revisión humana antes de tomar decisiones académicas."
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.save(output_path)


# ── Corrected document ─────────────────────────────────────────────────────────

def _add_markdown_content(doc: Document, text: str) -> None:
    """Parse lightweight markdown (# headings, - bullets) and add to document."""
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(stripped[2:])
        elif stripped:
            doc.add_paragraph(stripped)


def generate_corrected_docx(
    corrected_text: str,
    eval_filename: str,
    output_path: Path,
) -> None:
    """Create a Word document from the AI-corrected text."""
    doc = Document()

    for style_name in ("Normal", "List Bullet", "List Number"):
        try:
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style.font.size = Pt(11)
        except KeyError:
            pass

    now = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")

    # ── Cover ────────────────────────────────────────────────────────────────
    cover_title = doc.add_heading("DOCUMENTO CORREGIDO", 0)
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Versión con correcciones aplicadas · Claude Opus 4").italic = True

    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.add_run("Documento original: ").bold = True
    meta.add_run(eval_filename)

    meta2 = doc.add_paragraph()
    meta2.add_run("Fecha de corrección: ").bold = True
    meta2.add_run(now)

    doc.add_paragraph()
    notice = doc.add_paragraph()
    notice_run = notice.add_run(
        "AVISO: Este documento fue generado automáticamente por IA con las correcciones "
        "sugeridas por el evaluador. Revise el contenido antes de su uso definitivo."
    )
    notice_run.italic = True
    notice_run.font.size = Pt(9)
    notice_run.font.color.rgb = _AMBER

    doc.add_page_break()

    # ── Corrected content ────────────────────────────────────────────────────
    _add_markdown_content(doc, corrected_text)

    doc.save(output_path)
